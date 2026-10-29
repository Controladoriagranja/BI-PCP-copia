"""Importa a primeira aba de metas_produtividade.xlsx para pcp.metas, sem DDL."""
from __future__ import annotations

import argparse
import logging
import os
import shutil
import tempfile
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from produtividade_config import ROOT, carregar_env, conectar

TELA = 'calendario_producao_produtividade'
DIAS = [('segunda', 'SEG'), ('terca', 'TER'), ('quarta', 'QUA'), ('quinta', 'QUI'), ('sexta', 'SEX'), ('sabado', 'SAB')]
SERIES = [('AN_T1', 'Ave Nova 1º turno', 'diario_industria_turnos'),
          ('AN_T2', 'Ave Nova 2º turno', 'diario_industria_turnos'),
          ('RA', 'Real Alimentos', 'diario_industria')]
CODIGOS = tuple(f'PROD_{serie}_{dia}' for serie, _, _ in SERIES for _, dia in DIAS)
COLUNAS = {'id', 'indicador_codigo', 'indicador_nome', 'tela', 'tabela_origem', 'coluna_origem',
           'valor_meta', 'unidade', 'comparacao', 'referencia_mes', 'observacao', 'ativo', 'criado_em', 'atualizado_em'}


def normalizar(valor) -> str:
    texto = unicodedata.normalize('NFKD', str(valor or '')).encode('ascii', 'ignore').decode().lower()
    return ''.join(c for c in texto if c.isalnum())


def validar_numero(valor, celula: str) -> Decimal | None:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool) or not isinstance(valor, (int, float, Decimal)):
        raise ValueError(f'{celula}: informe número não negativo ou deixe vazio; fórmulas e textos não são aceitos.')
    try:
        numero = Decimal(str(valor))
    except InvalidOperation:
        raise ValueError(f'{celula}: número inválido.') from None
    if not numero.is_finite() or numero < 0:
        raise ValueError(f'{celula}: número deve ser finito e não negativo.')
    return numero


def ler_planilha(arquivo: Path) -> dict:
    from openpyxl import load_workbook
    # A validação inteira termina antes de abrir a conexão de escrita.
    with tempfile.TemporaryDirectory(prefix='pcp_metas_') as pasta:
        copia = Path(pasta) / 'metas.xlsx'
        shutil.copy2(arquivo, copia)
        wb = load_workbook(copia, read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb.worksheets[0]
            cabecalhos = tuple(normalizar(ws.cell(1, coluna).value) for coluna in (2, 3, 4))
            if cabecalhos != ('1t', '2t', 'realalimentos'):
                raise ValueError('Cabeçalhos B1:D1 devem ser 1T, 2T e Real Alimentos.')
            por_dia = {}
            aliases = {normalizar(nome + sufixo): codigo for nome, codigo in DIAS for sufixo in ('', '-feira')}
            for numero_linha, linha in enumerate(ws.iter_rows(min_row=2, max_col=4), start=2):
                valores = [celula.value for celula in linha]
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in valores):
                    continue
                dia = aliases.get(normalizar(valores[0]))
                if not dia or dia in por_dia:
                    raise ValueError(f'Linha {numero_linha}: dia ausente, repetido ou diferente de segunda a sábado.')
                por_dia[dia] = [validar_numero(c.value, f'{coluna}{numero_linha}') for coluna, c in zip('BCD', linha[1:])]
            if set(por_dia) != {codigo for _, codigo in DIAS}:
                raise ValueError('A planilha precisa conter os seis dias, de segunda a sábado.')
            return {f'PROD_{serie}_{dia}': por_dia[dia][indice]
                    for indice, (serie, _, _) in enumerate(SERIES) for _, dia in DIAS}
        finally:
            wb.close()


def verificar_estrutura(cur) -> bool:
    cur.execute('SELECT current_database()')
    if cur.fetchone()[0] != 'bi_granja':
        raise ValueError('Banco diferente de bi_granja; nenhuma gravação permitida.')
    cur.execute("SELECT column_name FROM information_schema.columns WHERE table_schema='pcp' AND table_name='metas'")
    faltantes = COLUNAS - {r[0] for r in cur.fetchall()}
    if faltantes:
        raise ValueError('pcp.metas divergente; faltam colunas: ' + ', '.join(sorted(faltantes)))
    cur.execute("""SELECT is_identity,column_default FROM information_schema.columns
        WHERE table_schema='pcp' AND table_name='metas' AND column_name='id'""")
    identidade, padrao = cur.fetchone()
    if identidade != 'YES' and not padrao:
        raise ValueError('pcp.metas.id sem geração automática; nenhuma alteração estrutural ou gravação realizada.')
    cur.execute("""SELECT EXISTS (
        SELECT 1 FROM pg_index i
        WHERE i.indrelid='pcp.metas'::regclass AND i.indisunique AND i.indisvalid
          AND i.indpred IS NULL AND i.indexprs IS NULL AND i.indnkeyatts=2
          AND (SELECT array_agg(a.attname::text ORDER BY a.attname::text)
               FROM unnest(i.indkey) WITH ORDINALITY k(attnum,ord)
               JOIN pg_attribute a ON a.attrelid=i.indrelid AND a.attnum=k.attnum
               WHERE k.ord<=i.indnkeyatts)=ARRAY['indicador_codigo','referencia_mes']::text[]
    )""")
    return bool(cur.fetchone()[0])


def metadados(codigo: str) -> tuple:
    for serie, nome, tabela in SERIES:
        for dia, sufixo in DIAS:
            if codigo == f'PROD_{serie}_{sufixo}':
                return (f'Produtividade {nome} - {dia}', tabela)
    raise ValueError('Código fora dos indicadores de Produtividade autorizados.')


def gravar(conn, valores: dict, hoje: date) -> dict:
    if set(valores) != set(CODIGOS):
        raise ValueError('Conjunto de indicadores diferente da planilha validada.')
    primeiro_atual = hoje.replace(day=1)
    resultado = {'inseridas': 0, 'atualizadas': 0, 'desativadas': 0, 'preservadas': 0}
    with conn.transaction():
        with conn.cursor() as cur:
            resultado['indice_unico'] = verificar_estrutura(cur)
            # Serializa update+insert mesmo quando a tabela não possui índice único.
            # Não cria/renomeia índices nem altera permissões, schema ou outros indicadores.
            cur.execute('LOCK TABLE pcp.metas IN SHARE ROW EXCLUSIVE MODE')
            cur.execute("""SELECT indicador_codigo,referencia_mes,count(*) FROM pcp.metas
                WHERE indicador_codigo=ANY(%s) GROUP BY indicador_codigo,referencia_mes HAVING count(*)>1""", (list(CODIGOS),))
            if cur.fetchall():
                raise ValueError('Metas PROD_ duplicadas por código/mês; transação cancelada sem corrigir dados automaticamente.')
            cur.execute('SELECT indicador_codigo,referencia_mes,tela FROM pcp.metas WHERE indicador_codigo=ANY(%s)', (list(CODIGOS),))
            existentes = cur.fetchall()
            if any(tela != TELA for _, _, tela in existentes):
                raise ValueError('Código PROD_ já utilizado por outra tela; nenhuma linha foi alterada.')
            inicial = not existentes
            resultado['carga_inicial'] = inicial
            meses = [date(ano, mes, 1) for ano in range(2026, hoje.year + 1) for mes in range(1, 13)
                     if inicial or date(ano, mes, 1) >= primeiro_atual]
            meses = sorted(set(meses) | {mes for _, mes, _ in existentes if mes >= primeiro_atual})
            resultado['preservadas'] = sum(1 for _, mes, _ in existentes if mes < primeiro_atual)
            for mes in meses:
                for codigo in CODIGOS:
                    valor = valores[codigo]
                    nome, tabela = metadados(codigo)
                    if valor is None:
                        cur.execute("""UPDATE pcp.metas SET ativo=false, atualizado_em=CURRENT_TIMESTAMP
                            WHERE indicador_codigo=%s AND referencia_mes=%s AND tela=%s AND ativo IS DISTINCT FROM false""",
                                    (codigo, mes, TELA))
                        resultado['desativadas'] += cur.rowcount
                        continue
                    payload = (nome, TELA, tabela, 'quantidade', valor, 'aves', '>=')
                    cur.execute("""UPDATE pcp.metas SET indicador_nome=%s,tela=%s,tabela_origem=%s,
                        coluna_origem=%s,valor_meta=%s,unidade=%s,comparacao=%s,observacao=NULL,ativo=true,
                        atualizado_em=CURRENT_TIMESTAMP WHERE indicador_codigo=%s AND referencia_mes=%s AND tela=%s
                        AND (indicador_nome,tela,tabela_origem,coluna_origem,valor_meta,unidade,comparacao,observacao,ativo)
                        IS DISTINCT FROM (%s,%s,%s,%s,%s,%s,%s,NULL,true)""",
                                payload + (codigo, mes, TELA) + payload)
                    resultado['atualizadas'] += cur.rowcount
                    cur.execute("""INSERT INTO pcp.metas (indicador_codigo,indicador_nome,tela,tabela_origem,
                        coluna_origem,valor_meta,unidade,comparacao,referencia_mes,observacao,ativo,criado_em,atualizado_em)
                        SELECT %s,%s,%s,%s,%s,%s,%s,%s,%s,NULL,true,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP
                        WHERE NOT EXISTS (SELECT 1 FROM pcp.metas WHERE indicador_codigo=%s AND referencia_mes=%s)""",
                                (codigo,) + payload + (mes, codigo, mes))
                    resultado['inseridas'] += cur.rowcount
    return resultado


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--env', type=Path)
    parser.add_argument('--validar', action='store_true', help='Valida apenas a planilha; não conecta ao banco.')
    args = parser.parse_args(argv)
    erro_config = None
    try:
        carregar_env(args.env)
    except Exception as exc:
        erro_config = exc
    caminho_log = Path(os.environ.get('METAS_PRODUTIVIDADE_LOG', ROOT / 'outputs/metas-produtividade/robo.log'))
    caminho_log.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(filename=caminho_log, level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s')
    try:
        if erro_config is not None:
            raise erro_config
        arquivo = os.environ.get('METAS_PRODUTIVIDADE_XLSX')
        if not arquivo:
            raise ValueError('Defina METAS_PRODUTIVIDADE_XLSX na configuração.')
        valores = ler_planilha(Path(arquivo))
        if args.validar:
            logging.info('Planilha validada; %d metas preenchidas; nenhuma conexão SQL.', sum(v is not None for v in valores.values()))
            print('Planilha válida. Nenhuma gravação realizada.')
            return 0
        with conectar() as conn:
            resultado = gravar(conn, valores, datetime.now(ZoneInfo('America/Sao_Paulo')).date())
        logging.info('Carga concluída: %s', resultado)
        print(resultado)
        return 0
    except Exception as exc:
        # Não escreve DSN, credenciais ou mensagens do driver no log.
        mensagem = str(exc) if isinstance(exc, (ValueError, FileNotFoundError, PermissionError)) else type(exc).__name__
        logging.error('Carga cancelada: %s', mensagem)
        print('Carga cancelada. Consulte o log; nenhuma carga parcial foi mantida.')
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
